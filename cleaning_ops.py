"""
cleaning_ops.py — structured, human-approved workbook operations
=====================================================================
The Data Cleaning Agent can only PROPOSE operations (cleaning_store.create_operation,
status='proposed'). Everything in this module that actually writes to a workbook
(apply_operation) is called from a Flask route triggered by a person clicking
Approve — never from an agent tool. This is the enforcement point for "no silent
modification": even if the LLM decided to skip asking, nothing changes until a
human calls approve_operation().

Workbooks are edited in place with openpyxl (formulas, styles, merged cells,
frozen panes, etc. are preserved by construction — we only ever touch specific
cells, never rebuild the sheet). CSV files are edited with the csv module for the
same reason (no pandas round-trip that could reformat values).

Every applied cell change is recorded as {sheet, row, column, before, after} so
undo can restore the exact prior value.
"""

import csv
import io
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl
import pandas as pd

import cleaning_store as store

ALLOWED_OPERATIONS = {
    "replace_value",
    "trim_whitespace",
    "normalize_category",
    "fill_missing",
    "change_data_type",
    "correct_date",
    "flag_record",
}

# remove_duplicate is intentionally implemented as flag_record: deleting rows in
# a workbook in place breaks row-relative formulas and named ranges, so we flag
# instead and let the person delete manually if they choose to.
REMOVE_DUPLICATE_ALIAS = "flag_record"

_EMPTY_SENTINELS = (None, "", "nan", "NaN", "NaT")


class OperationError(ValueError):
    pass


def _is_empty(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and pd.isna(v):
        return True
    if isinstance(v, str) and v.strip() == "":
        return True
    return False


def _values_equal(a: Any, b: Any) -> bool:
    """Used to check a cell hasn't drifted since an operation was proposed. Exact
    for strings — do NOT strip whitespace here, or a stale-value check would be
    blind to exactly the whitespace edits this tool exists to make."""
    if _is_empty(a) and _is_empty(b):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < 1e-9
    return str(a) == str(b)


def _load_workbook(path: Path, keep_vba: bool = False):
    return openpyxl.load_workbook(path, data_only=False, keep_vba=keep_vba)


def _header_map(ws) -> dict[str, int]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    return {str(c.value): c.column for c in header_row if c.value is not None}


def _get_or_create_column(ws, header_map: dict[str, int], name: str) -> int:
    if name in header_map:
        return header_map[name]
    new_col = ws.max_column + 1
    ws.cell(row=1, column=new_col).value = name  # DC-01: ws.cell(value=X) silently no-ops when X is None
    header_map[name] = new_col
    return new_col


# ─────────────────────────────────────────────
# Validation (re-checked at approval time, not just at proposal time)
# ─────────────────────────────────────────────

CSV_ALLOWED_OPERATIONS = ("replace_value", "trim_whitespace", "normalize_category",
                           "fill_missing", "change_data_type", "correct_date", "flag_record")


def validate_operation(op: dict, session: dict) -> list[str]:
    """Raises OperationError for anything that would make the operation unsafe to
    apply; returns a list of non-fatal warnings otherwise. Re-checked at approval
    time (not just proposal time), including a stale-value check — the current
    cell must still equal the proposal's recorded old_value — for BOTH xlsx and
    csv files (the csv path previously skipped this check entirely, DC-04)."""
    operation = op.get("operation")
    if operation not in ALLOWED_OPERATIONS:
        raise OperationError(f"'{operation}' is not an allowed operation")

    file_type = session["file_type"]
    working_path = Path(session["working_path"])
    if not working_path.exists():
        raise OperationError("working file is missing")

    sheet = op.get("sheet") or session.get("sheet")
    if file_type == "csv":
        if operation not in CSV_ALLOWED_OPERATIONS:
            raise OperationError(f"'{operation}' is not supported for CSV files")
        return _validate_csv(op, working_path)
    return _validate_xlsx(op, working_path, sheet)


def _validate_xlsx(op: dict, working_path: Path, sheet: str) -> list[str]:
    warnings: list[str] = []
    wb = _load_workbook(working_path)
    if sheet not in wb.sheetnames:
        raise OperationError(f"sheet '{sheet}' does not exist in the working file")
    ws = wb[sheet]
    header_map = _header_map(ws)
    operation = op.get("operation")

    if operation == "flag_record":
        flag_name = op.get("flag_name")
        if not flag_name:
            raise OperationError("flag_record requires flag_name")
        targets = op.get("targets", [])
        if not targets:
            raise OperationError("flag_record requires at least one target row")
        for t in targets:
            row = t.get("row")
            if not isinstance(row, int) or row < 2 or row > ws.max_row:
                raise OperationError(f"row {row} is out of range for sheet '{sheet}'")
        return warnings

    column = op.get("column")
    if not column:
        raise OperationError(f"'{operation}' requires a column")
    if column not in header_map:
        raise OperationError(f"column '{column}' does not exist in sheet '{sheet}'")
    col_idx = header_map[column]

    targets = op.get("targets", [])
    if not targets:
        raise OperationError(f"'{operation}' requires at least one target")

    for t in targets:
        row = t.get("row")
        if not isinstance(row, int) or row < 2 or row > ws.max_row:
            raise OperationError(f"row {row} is out of range for sheet '{sheet}'")
        cell = ws.cell(row=row, column=col_idx)
        if cell.data_type == "f":
            raise OperationError(f"cell {column}{row} contains a formula and cannot be edited")
        current = cell.value
        old_value = t.get("old_value")
        if not _values_equal(current, old_value):
            raise OperationError(
                f"cell {column}{row} has changed since this operation was proposed "
                f"(expected {old_value!r}, found {current!r}) — refresh and re-propose"
            )
    return warnings


def _validate_csv(op: dict, working_path: Path) -> list[str]:
    warnings: list[str] = []
    with open(working_path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows:
        raise OperationError("CSV file is empty")
    header = rows[0]
    operation = op.get("operation")
    max_row = len(rows)  # rows[0] is the header, i.e. Excel row 1

    if operation == "flag_record":
        flag_name = op.get("flag_name")
        if not flag_name:
            raise OperationError("flag_record requires flag_name")
        targets = op.get("targets", [])
        if not targets:
            raise OperationError("flag_record requires at least one target row")
        for t in targets:
            row = t.get("row")
            if not isinstance(row, int) or row < 2 or row > max_row:
                raise OperationError(f"row {row} is out of range")
        return warnings

    column = op.get("column")
    if not column:
        raise OperationError(f"'{operation}' requires a column")
    if column not in header:
        raise OperationError(f"column '{column}' does not exist")
    col_idx = header.index(column)

    targets = op.get("targets", [])
    if not targets:
        raise OperationError(f"'{operation}' requires at least one target")

    for t in targets:
        row = t.get("row")
        if not isinstance(row, int) or row < 2 or row > max_row:
            raise OperationError(f"row {row} is out of range")
        data_row = rows[row - 1]
        current = data_row[col_idx] if col_idx < len(data_row) else None
        old_value = t.get("old_value")
        if not _values_equal(current, old_value):
            raise OperationError(
                f"cell {column}{row} has changed since this operation was proposed "
                f"(expected {old_value!r}, found {current!r}) — refresh and re-propose"
            )
    return warnings


# ─────────────────────────────────────────────
# Dry run
# ─────────────────────────────────────────────

def dry_run_operation(op: dict, session: dict) -> dict:
    warnings = validate_operation(op, session)
    preview = {"operation": op.get("operation"), "sheet": op.get("sheet") or session.get("sheet"),
               "warnings": warnings, "changes": []}

    if op.get("operation") == "flag_record":
        for t in op.get("targets", []):
            preview["changes"].append({
                "row": t["row"], "column": op.get("flag_name"),
                "before": None, "after": op.get("flag_value", "YES"),
            })
        return preview

    column = op.get("column")
    for t in op.get("targets", []):
        preview["changes"].append({
            "row": t["row"], "column": column,
            "before": t.get("old_value"), "after": t.get("new_value"),
        })
    return preview


# ─────────────────────────────────────────────
# Apply (openpyxl, in place) — only ever called after a human approval
# ─────────────────────────────────────────────

def _apply_xlsx(op: dict, session: dict) -> list[dict]:
    working_path = Path(session["working_path"])
    keep_vba = session["file_type"] == "xlsm"
    wb = _load_workbook(working_path, keep_vba=keep_vba)
    sheet = op.get("sheet") or session.get("sheet")
    ws = wb[sheet]
    header_map = _header_map(ws)
    cell_changes = []

    if op["operation"] == "flag_record":
        flag_name = op.get("flag_name", "DQ_FLAG")
        flag_value = op.get("flag_value", "YES")
        col_idx = _get_or_create_column(ws, header_map, flag_name)
        for t in op.get("targets", []):
            row = t["row"]
            cell = ws.cell(row=row, column=col_idx)
            before = cell.value
            cell.value = flag_value
            cell_changes.append({"sheet": sheet, "row": row, "column": flag_name,
                                  "before": before, "after": flag_value})
    else:
        column = op["column"]
        col_idx = header_map[column]
        for t in op.get("targets", []):
            row = t["row"]
            cell = ws.cell(row=row, column=col_idx)
            before = cell.value
            new_value = t.get("new_value")
            cell.value = new_value
            cell_changes.append({"sheet": sheet, "row": row, "column": column,
                                  "before": before, "after": new_value})

    wb.save(working_path)
    return cell_changes


def _apply_csv(op: dict, session: dict) -> list[dict]:
    working_path = Path(session["working_path"])
    with open(working_path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    if not rows:
        raise OperationError("CSV file is empty")
    header = rows[0]
    cell_changes = []

    if op["operation"] == "flag_record":
        flag_name = op.get("flag_name", "DQ_FLAG")
        flag_value = op.get("flag_value", "YES")
        if flag_name in header:
            col_idx = header.index(flag_name)
        else:
            col_idx = len(header)
            header.append(flag_name)
            for r in rows[1:]:
                r.append("")
        for t in op.get("targets", []):
            row = t["row"]
            data_idx = row - 1  # rows[0] is header, data row `row` (excel, 1-based+header) -> rows[row-1]
            if data_idx >= len(rows):
                raise OperationError(f"row {row} is out of range")
            before = rows[data_idx][col_idx] if col_idx < len(rows[data_idx]) else None
            while len(rows[data_idx]) <= col_idx:
                rows[data_idx].append("")
            rows[data_idx][col_idx] = flag_value
            cell_changes.append({"sheet": "Sheet1", "row": row, "column": flag_name,
                                  "before": before, "after": flag_value})
    else:
        column = op["column"]
        if column not in header:
            raise OperationError(f"column '{column}' does not exist")
        col_idx = header.index(column)
        for t in op.get("targets", []):
            row = t["row"]
            data_idx = row - 1
            if data_idx >= len(rows):
                raise OperationError(f"row {row} is out of range")
            before = rows[data_idx][col_idx] if col_idx < len(rows[data_idx]) else None
            new_value = t.get("new_value")
            rows[data_idx][col_idx] = "" if new_value is None else str(new_value)
            cell_changes.append({"sheet": "Sheet1", "row": row, "column": column,
                                  "before": before, "after": new_value})

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(rows)
    working_path.write_text(buf.getvalue(), encoding="utf-8")
    return cell_changes


def apply_operation(op_row: dict, session: dict) -> dict:
    """op_row is a row from cleaning_store (has .op, .id). Only call this after a
    human has approved. Re-validates before writing, then records cell_changes."""
    op = op_row["op"]
    validate_operation(op, session)

    if session["file_type"] == "csv":
        cell_changes = _apply_csv(op, session)
    else:
        cell_changes = _apply_xlsx(op, session)

    now = datetime.now().isoformat()
    return store.update_operation(
        op_row["id"], status="applied", cell_changes=cell_changes, applied_at=now,
    )


# ─────────────────────────────────────────────
# Undo / revert
# ─────────────────────────────────────────────

def undo_operation(op_row: dict, session: dict) -> dict:
    if op_row["status"] != "applied":
        raise OperationError(f"operation {op_row['id']} is not currently applied (status={op_row['status']})")
    cell_changes = op_row["cell_changes"]

    if session["file_type"] == "csv":
        working_path = Path(session["working_path"])
        with open(working_path, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
        header = rows[0]
        for ch in reversed(cell_changes):
            if ch["column"] not in header:
                continue
            col_idx = header.index(ch["column"])
            data_idx = ch["row"] - 1
            if data_idx < len(rows) and col_idx < len(rows[data_idx]):
                rows[data_idx][col_idx] = "" if ch["before"] is None else str(ch["before"])
        buf = io.StringIO()
        csv.writer(buf).writerows(rows)
        working_path.write_text(buf.getvalue(), encoding="utf-8")
    else:
        working_path = Path(session["working_path"])
        keep_vba = session["file_type"] == "xlsm"
        wb = _load_workbook(working_path, keep_vba=keep_vba)
        for ch in reversed(cell_changes):
            ws = wb[ch["sheet"]]
            header_map = _header_map(ws)
            if ch["column"] not in header_map:
                continue
            # DC-01: ws.cell(..., value=X) silently no-ops when X is None (openpyxl only
            # assigns when value is not None), which would leave a cell that was originally
            # empty stuck at its post-apply value. Direct attribute assignment always writes,
            # including None, so undo can restore None/""/0/False exactly.
            ws.cell(row=ch["row"], column=header_map[ch["column"]]).value = ch["before"]
        wb.save(working_path)

    now = datetime.now().isoformat()
    return store.update_operation(op_row["id"], status="undone", undone_at=now)


def revert_session(session: dict) -> dict:
    """Discards all applied changes: restores working file from the untouched
    original and marks every applied operation as reverted."""
    original_path = Path(session["original_path"])
    working_path = Path(session["working_path"])
    working_path.write_bytes(original_path.read_bytes())
    for op_row in store.list_operations(session["id"], status="applied"):
        store.update_operation(op_row["id"], status="reverted", undone_at=datetime.now().isoformat())
    return store.update_session(session["id"], state="DETECT")


# ─────────────────────────────────────────────
# Validation after changes
# ─────────────────────────────────────────────

def summarize_findings(findings: list[dict]) -> dict:
    by_check: dict[str, int] = {}
    total = 0
    for f in findings:
        by_check[f["check_name"]] = by_check.get(f["check_name"], 0) + f["count"]
        total += f["count"]
    return {"total_findings": total, "by_check": by_check}


def compare_before_after(before_findings: list[dict], after_findings: list[dict]) -> dict:
    before = summarize_findings(before_findings)
    after = summarize_findings(after_findings)
    return {
        "before": before,
        "after": after,
        "resolved_count": max(0, before["total_findings"] - after["total_findings"]),
    }


# ─────────────────────────────────────────────
# Finalize: workbook is already edited in place, so finalizing is just handing
# back the working copy under its output name (no rebuild, no lost formatting)
# ─────────────────────────────────────────────

def finalize_workbook(session: dict) -> Path:
    working_path = Path(session["working_path"])
    cleaned_path, _ = store.output_paths(session["id"])
    cleaned_path = cleaned_path.with_suffix(working_path.suffix)
    cleaned_path.write_bytes(working_path.read_bytes())
    return cleaned_path


REPORT_WARNINGS = (
    "Charts, images, and pivot tables are not guaranteed to survive an openpyxl "
    "edit; check the cleaned workbook visually if the source contained any."
)


def generate_report(session: dict, before_findings: list[dict], after_findings: list[dict]) -> Path:
    groups = store.list_anomaly_groups(session["id"])
    operations = store.list_operations(session["id"])
    research = store.list_research(session["id"])
    comparison = compare_before_after(before_findings, after_findings)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Nexus Data Cleaning Report"])
    ws.append(["Session", session["id"]])
    ws.append(["Original file", session["original_filename"]])
    ws.append(["Sheet", session.get("sheet")])
    ws.append(["Generated", datetime.now().isoformat()])
    ws.append([])
    ws.append(["Findings before cleaning", comparison["before"]["total_findings"]])
    ws.append(["Findings after cleaning", comparison["after"]["total_findings"]])
    ws.append(["Resolved", comparison["resolved_count"]])
    ws.append([])
    ws.append(["Note", REPORT_WARNINGS])

    dims = wb.create_sheet("Readiness Dimensions")
    dims.append(["Dimension", "Evidence"])
    dim_notes: dict[str, list[str]] = {}
    for g in groups:
        dim_notes.setdefault(g["dimension"], []).append(f"{g['check_name']}: {g['count']} finding(s)")
    for dim, notes in dim_notes.items():
        dims.append([dim, "; ".join(notes)])

    ag = wb.create_sheet("Anomaly Groups")
    ag.append(["ID", "Check", "Dimension", "Columns", "Count", "Classification",
               "Detection Conf.", "Evidence Conf.", "Recommendation Conf.", "Status"])
    for g in groups:
        ag.append([g["id"], g["check_name"], g["dimension"], ", ".join(g["columns"]), g["count"],
                   g["classification"], g["detection_confidence"], g["evidence_confidence"],
                   g["recommendation_confidence"], g["status"]])

    def _ops_sheet(name: str, status: str):
        sh = wb.create_sheet(name)
        sh.append(["ID", "Operation", "Column/Sheet", "Reason", "Cell changes", "Status"])
        for op in operations:
            if op["status"] != status:
                continue
            sh.append([op["id"], op["op"].get("operation"), op["op"].get("column") or op["op"].get("sheet"),
                       op["reason"], len(op["cell_changes"]), op["status"]])

    _ops_sheet("Approved", "applied")
    _ops_sheet("Rejected", "rejected")

    unresolved = wb.create_sheet("Unresolved")
    unresolved.append(["ID", "Check", "Count", "Note"])
    for g in groups:
        if g["status"] == "open":
            unresolved.append([g["id"], g["check_name"], g["count"], g.get("evidence", {}).get("note", "")])

    hist = wb.create_sheet("Historical Evidence")
    hist.append(["Group ID", "Evidence"])
    for g in groups:
        if g.get("evidence"):
            hist.append([g["id"], str(g["evidence"])])

    ext = wb.create_sheet("External Research")
    ext.append(["Query", "Source", "URL", "Fact", "Confidence", "Retrieved"])
    for r in research:
        ext.append([r["query"], r["source_title"], r["url"], r["fact"], r["confidence"], r["retrieved_at"]])

    val = wb.create_sheet("Validation")
    val.append(["Metric", "Before", "After"])
    checks = set(comparison["before"]["by_check"]) | set(comparison["after"]["by_check"])
    for c in sorted(checks):
        val.append([c, comparison["before"]["by_check"].get(c, 0), comparison["after"]["by_check"].get(c, 0)])

    _, report_path = store.output_paths(session["id"])
    wb.save(report_path)
    return report_path
