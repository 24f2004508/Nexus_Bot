"""
cleaning_tools.py — deterministic data-quality tools for the Data Cleaning Agent
====================================================================================
Every function here is plain Python: profiling, workbook/CSV I/O, and detection
checks. Nothing here calls the LLM. The Data Cleaning Agent (cleaning_agent.py)
wraps a subset of these as tools so the LLM can call them, but the checks
themselves are deterministic and testable on their own.

Row references use the actual Excel row number (data starts at row 2, since row 1
is the header), never a pandas positional index — those go stale the moment a row
is dropped. Nothing in this module mutates a workbook; that happens in cleaning_ops.py.

data_contract shape (confirmed by the user via the agent, not inferred silently):
    {
        "mandatory_columns": [...],
        "dob_column": "...",
        "status_column": "...",
        "id_column": "...",
        "date_column": "...",           # used to order rows for status-transition checks
        "date_pairs": [[start_col, end_col], ...],
        "column_definitions": {col: "what this column means, confirmed by the user"},
    }
"""

import difflib
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl
import pandas as pd

# ─────────────────────────────────────────────
# Reused/ported from the old data_cleaning.py
# ─────────────────────────────────────────────

_DEFAULT_VALID_TRANSITIONS = {
    "in force": {"lapsed", "cancelled", "matured", "surrendered", "in force"},
    "lapsed": {"in force", "cancelled", "lapsed"},
    "cancelled": {"cancelled"},
    "matured": {"matured"},
    "surrendered": {"surrendered"},
    "pending": {"in force", "cancelled", "pending"},
}

AMOUNT_KEYWORDS = (
    "premium", "loss", "claim", "amount", "sum", "value", "exposure",
    "price", "cost", "fee", "charge", "income", "salary", "reserve", "payout",
)
ID_KEYWORDS = ("id", "no", "number", "ref")


def _is_orange(rgb_hex: Optional[str]) -> bool:
    if not rgb_hex or len(rgb_hex) < 6:
        return False
    try:
        r, g, b = int(rgb_hex[-6:-4], 16), int(rgb_hex[-4:-2], 16), int(rgb_hex[-2:], 16)
    except ValueError:
        return False
    return r >= 180 and b <= 120 and b < g < r


def _cell_fill_rgb(cell) -> Optional[str]:
    fill = cell.fill
    if fill is None or fill.fgColor is None:
        return None
    color = fill.fgColor
    if getattr(color, "type", None) == "rgb" and isinstance(color.rgb, str):
        return color.rgb
    return None


def _parse_dob(val) -> Optional[pd.Timestamp]:
    if pd.isna(val):
        return None
    ts = pd.to_datetime(val, dayfirst=True, errors="coerce")
    return None if pd.isna(ts) else ts


def _normalize(val: Any) -> str:
    return re.sub(r"\s+", " ", str(val).strip().lower())


def _is_text_column(series: pd.Series) -> bool:
    """pandas >= 3.0 defaults inferred string columns to a 'str' dtype, not the
    legacy numpy `object` dtype — check both so text columns aren't silently skipped."""
    return series.dtype == object or pd.api.types.is_string_dtype(series)


# ─────────────────────────────────────────────
# File / workbook I/O
# ─────────────────────────────────────────────

SUPPORTED_EXTENSIONS = {"csv", "xlsx", "xlsm"}


def list_sheets(path: Path, file_type: str) -> list[str]:
    if file_type == "csv":
        return ["Sheet1"]
    xl = pd.ExcelFile(path)
    return list(xl.sheet_names)


def read_sheet_df(path: Path, file_type: str, sheet: str) -> pd.DataFrame:
    if file_type == "csv":
        return pd.read_csv(path)
    return pd.read_excel(path, sheet_name=sheet)


def list_columns(path: Path, file_type: str, sheet: str) -> list[str]:
    df = read_sheet_df(path, file_type, sheet)
    return list(df.columns.astype(str))


def check_ragged_rows(df: pd.DataFrame) -> Optional[str]:
    """CSV integrity check: a row with an unescaped/unquoted comma in a text
    field (more fields than the header) makes pandas' lenient parser shift
    that row's extra field into the DataFrame's index instead of raising —
    the header and data silently go out of alignment (a column can appear to
    vanish, and every later row-number-based lookup in cleaning_tools.py
    assumes a plain 0-based position, which such a shift breaks). Detected by
    checking whether the DataFrame still has the plain positional index every
    other read produces; used at upload time to reject the file with a clear
    message instead of silently analysing misaligned data or crashing later
    deep in a detection check."""
    if not isinstance(df.index, pd.RangeIndex):
        return (
            "the file's rows don't line up consistently with its header — this usually means "
            "a text field contains a comma that isn't wrapped in quotes. Please quote any text "
            "fields that contain commas (most spreadsheet tools do this automatically when you "
            "export CSV) and re-upload."
        )
    return None


# ─────────────────────────────────────────────
# Profiling
# ─────────────────────────────────────────────

def profile_dataframe(df: pd.DataFrame, columns: Optional[list[str]] = None) -> dict:
    cols = columns or list(df.columns)
    n_rows = len(df)
    out = {"n_rows": n_rows, "n_cols": len(cols), "columns": []}
    for col in cols:
        if col not in df.columns:
            continue
        series = df[col]
        null_count = int(series.isna().sum())
        entry = {
            "name": col,
            "dtype": str(series.dtype),
            "null_count": null_count,
            "null_pct": round(100 * null_count / n_rows, 2) if n_rows else 0.0,
        }
        numeric = pd.to_numeric(series, errors="coerce")
        non_null_numeric = numeric.dropna()
        if len(non_null_numeric) > 0 and len(non_null_numeric) >= 0.5 * series.dropna().shape[0]:
            q1, q3 = non_null_numeric.quantile(0.25), non_null_numeric.quantile(0.75)
            iqr = q3 - q1
            lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
            entry.update({
                "mean": round(float(non_null_numeric.mean()), 4),
                "min": float(non_null_numeric.min()),
                "max": float(non_null_numeric.max()),
                "outlier_count": int(((non_null_numeric < lo) | (non_null_numeric > hi)).sum()),
            })
        else:
            vc = series.dropna().astype(str).value_counts().head(5)
            entry.update({
                "top_values": [{"value": k, "count": int(v)} for k, v in vc.items()],
                "unique_count": int(series.dropna().astype(str).nunique()),
            })
        out["columns"].append(entry)
    return out


# ─────────────────────────────────────────────
# Findings (each check returns 0 or more finding dicts, later stored as anomaly_groups)
# A finding: {check_name, dimension, columns, row_refs, examples, count, evidence}
# row_refs: list of {"row": <excel row>, "column": <name or None>}
# ─────────────────────────────────────────────

def detect_missing_values(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    findings = []
    for col in columns:
        if col not in df.columns:
            continue
        series = df[col]
        is_missing = series.isna() | series.astype(str).str.strip().eq("")
        idx = series.index[is_missing]
        if len(idx) == 0:
            continue
        findings.append({
            "check_name": "missing_values",
            "dimension": "Completeness",
            "columns": [col],
            "row_refs": [{"row": int(i) + 2, "column": col} for i in idx],
            "examples": [],
            "count": int(len(idx)),
            "evidence": {"note": "null, blank, or whitespace-only values"},
        })
    return findings


def detect_duplicate_rows(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    subset = [c for c in columns if c in df.columns] or None
    dup_mask = df.duplicated(subset=subset, keep=False)
    idx = df.index[dup_mask]
    if len(idx) == 0:
        return []
    return [{
        "check_name": "duplicate_rows",
        "dimension": "Consistency",
        "columns": subset or list(df.columns),
        "row_refs": [{"row": int(i) + 2, "column": None} for i in idx],
        "examples": [],
        "count": int(len(idx)),
        "evidence": {"note": "exact duplicate rows across the selected columns"},
    }]


def detect_duplicate_ids(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    findings = []
    for col in columns:
        if col not in df.columns:
            continue
        if not any(kw in col.lower() for kw in ID_KEYWORDS):
            continue
        non_null = df[col].dropna()
        dup_vals = non_null[non_null.duplicated(keep=False)]
        if len(dup_vals) == 0:
            continue
        findings.append({
            "check_name": "duplicate_ids",
            "dimension": "Referential Integrity",
            "columns": [col],
            "row_refs": [{"row": int(i) + 2, "column": col} for i in dup_vals.index],
            "examples": [str(v) for v in dup_vals.unique()[:5]],
            "count": int(len(dup_vals)),
            "evidence": {"note": f"column name suggests an identifier: '{col}'"},
        })
    return findings


def detect_type_inconsistency(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    findings = []
    for col in columns:
        if col not in df.columns or not _is_text_column(df[col]):
            continue
        non_null = df[col].dropna().astype(str)
        if len(non_null) == 0:
            continue
        numeric = pd.to_numeric(non_null, errors="coerce")
        numeric_like = numeric.notna()
        frac_numeric = numeric_like.mean()
        if 0 < frac_numeric < 1:
            bad_idx = non_null.index[~numeric_like]
            findings.append({
                "check_name": "type_inconsistency",
                "dimension": "Validity",
                "columns": [col],
                "row_refs": [{"row": int(i) + 2, "column": col} for i in bad_idx],
                "examples": [str(v) for v in non_null.loc[bad_idx].unique()[:5]],
                "count": int(len(bad_idx)),
                "evidence": {"note": f"{round(100*frac_numeric,1)}% of values in '{col}' parse as numbers; these don't"},
            })
    return findings


def detect_whitespace_case_variants(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    findings = []
    for col in columns:
        if col not in df.columns or not _is_text_column(df[col]):
            continue
        series = df[col].dropna().astype(str)
        if len(series) == 0:
            continue
        has_ws = series != series.str.strip()
        ws_idx = series.index[has_ws]
        if len(ws_idx) > 0:
            unique_bad = list(series.loc[ws_idx].unique()[:5])
            findings.append({
                "check_name": "whitespace",
                "dimension": "Consistency",
                "columns": [col],
                "row_refs": [{"row": int(i) + 2, "column": col} for i in ws_idx],
                "examples": unique_bad,
                "count": int(len(ws_idx)),
                "evidence": {
                    "note": "leading/trailing whitespace",
                    # exact strings to use verbatim as old_value/new_value — do not
                    # retype or re-quote these when proposing an operation
                    "suggested_fixes": {v: v.strip() for v in unique_bad},
                },
            })
        groups: dict[str, set[str]] = {}
        for v in series.unique():
            key = v.strip().lower()
            groups.setdefault(key, set()).add(v)
        case_variant_keys = {k: vs for k, vs in groups.items() if len(vs) > 1}
        if case_variant_keys:
            affected = set()
            for vs in case_variant_keys.values():
                affected |= vs
            idx = series.index[series.isin(affected)]
            findings.append({
                "check_name": "case_variants",
                "dimension": "Consistency",
                "columns": [col],
                "row_refs": [{"row": int(i) + 2, "column": col} for i in idx],
                "examples": [sorted(vs) for vs in list(case_variant_keys.values())[:5]],
                "count": int(len(idx)),
                "evidence": {"note": "same value in different capitalization"},
            })
    return findings


def detect_near_spelling_variants(df: pd.DataFrame, columns: list[str], cutoff: float = 0.82,
                                   max_cardinality: int = 200) -> list[dict]:
    findings = []
    for col in columns:
        if col not in df.columns or not _is_text_column(df[col]):
            continue
        distinct = sorted(set(str(v).strip() for v in df[col].dropna().unique()))
        if not (1 < len(distinct) <= max_cardinality):
            continue
        clusters: list[list[str]] = []
        used = set()
        for v in distinct:
            if v in used:
                continue
            close = difflib.get_close_matches(v, [d for d in distinct if d not in used], n=5, cutoff=cutoff)
            close = [c for c in close if c != v]
            if close:
                cluster = [v] + close
                clusters.append(cluster)
                used.update(cluster)
        if not clusters:
            continue
        affected_values = {v for cluster in clusters for v in cluster}
        idx = df.index[df[col].astype(str).str.strip().isin(affected_values)]
        findings.append({
            "check_name": "near_spelling_variants",
            "dimension": "Consistency",
            "columns": [col],
            "row_refs": [{"row": int(i) + 2, "column": col} for i in idx],
            "examples": clusters[:5],
            "count": int(len(idx)),
            "evidence": {"note": "textually similar values that may represent the same category"},
        })
    return findings


def detect_outliers(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    findings = []
    for col in columns:
        if col not in df.columns:
            continue
        numeric = pd.to_numeric(df[col], errors="coerce")
        non_null = numeric.dropna()
        if len(non_null) < 5:
            continue
        q1, q3 = non_null.quantile(0.25), non_null.quantile(0.75)
        iqr = q3 - q1
        if iqr == 0:
            continue
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        mask = (numeric < lo) | (numeric > hi)
        idx = numeric.index[mask.fillna(False)]
        if len(idx) == 0:
            continue
        findings.append({
            "check_name": "statistical_outliers",
            "dimension": "Validity",
            "columns": [col],
            "row_refs": [{"row": int(i) + 2, "column": col} for i in idx],
            "examples": [float(numeric.loc[i]) for i in list(idx)[:5]],
            "count": int(len(idx)),
            "evidence": {"note": f"outside 1.5x IQR [{lo:.2f}, {hi:.2f}] — potential anomaly, not automatically an error"},
        })
    return findings


def detect_negative_zero_amounts(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    findings = []
    for col in columns:
        if col not in df.columns or not any(kw in col.lower() for kw in AMOUNT_KEYWORDS):
            continue
        numeric = pd.to_numeric(df[col], errors="coerce")
        mask = numeric <= 0
        idx = numeric.index[mask.fillna(False)]
        if len(idx) == 0:
            continue
        findings.append({
            "check_name": "negative_or_zero_amount",
            "dimension": "Validity",
            "columns": [col],
            "row_refs": [{"row": int(i) + 2, "column": col} for i in idx],
            "examples": [float(numeric.loc[i]) for i in list(idx)[:5]],
            "count": int(len(idx)),
            "evidence": {"note": "negative/zero values in an amount-like column — could be refunds/reversals, not necessarily errors"},
        })
    return findings


def detect_invalid_future_dates(df: pd.DataFrame, columns: list[str]) -> list[dict]:
    findings = []
    today = pd.Timestamp(datetime.now().date())
    for col in columns:
        if col not in df.columns:
            continue
        lower = col.lower()
        if "date" not in lower and "dob" not in lower:
            continue
        non_null = df[col].dropna()
        if len(non_null) == 0:
            continue
        parsed = pd.to_datetime(non_null, errors="coerce", dayfirst=True)
        invalid_idx = non_null.index[parsed.isna()]
        if len(invalid_idx) > 0:
            findings.append({
                "check_name": "invalid_date",
                "dimension": "Temporal Integrity",
                "columns": [col],
                "row_refs": [{"row": int(i) + 2, "column": col} for i in invalid_idx],
                "examples": [str(v) for v in non_null.loc[invalid_idx].unique()[:5]],
                "count": int(len(invalid_idx)),
                "evidence": {"note": "value does not parse as a date"},
            })
        future_idx = parsed.index[(parsed > today).fillna(False)]
        if len(future_idx) > 0:
            findings.append({
                "check_name": "future_date",
                "dimension": "Temporal Integrity",
                "columns": [col],
                "row_refs": [{"row": int(i) + 2, "column": col} for i in future_idx],
                "examples": [str(parsed.loc[i].date()) for i in list(future_idx)[:5]],
                "count": int(len(future_idx)),
                "evidence": {"note": "date is after today"},
            })
    return findings


def detect_date_pair_order(df: pd.DataFrame, data_contract: dict) -> list[dict]:
    """Only runs for date pairs the user has explicitly confirmed in data_contract['date_pairs']."""
    findings = []
    for pair in data_contract.get("date_pairs", []):
        if len(pair) != 2:
            continue
        start_col, end_col = pair
        if start_col not in df.columns or end_col not in df.columns:
            continue
        start = pd.to_datetime(df[start_col], errors="coerce", dayfirst=True)
        end = pd.to_datetime(df[end_col], errors="coerce", dayfirst=True)
        mask = (start.notna()) & (end.notna()) & (start > end)
        idx = df.index[mask]
        if len(idx) == 0:
            continue
        findings.append({
            "check_name": "date_pair_order",
            "dimension": "Temporal Integrity",
            "columns": [start_col, end_col],
            "row_refs": [{"row": int(i) + 2, "column": start_col} for i in idx],
            "examples": [],
            "count": int(len(idx)),
            "evidence": {"note": f"'{start_col}' is after '{end_col}', confirmed as a start/end pair by the user"},
        })
    return findings


def detect_missing_mandatory(df: pd.DataFrame, data_contract: dict) -> list[dict]:
    findings = []
    for col in data_contract.get("mandatory_columns", []):
        if col not in df.columns:
            continue
        series = df[col]
        is_missing = series.isna() | series.astype(str).str.strip().eq("")
        idx = series.index[is_missing]
        if len(idx) == 0:
            continue
        findings.append({
            "check_name": "missing_mandatory_field",
            "dimension": "Completeness",
            "columns": [col],
            "row_refs": [{"row": int(i) + 2, "column": col} for i in idx],
            "examples": [],
            "count": int(len(idx)),
            "evidence": {"note": f"'{col}' is a confirmed mandatory field"},
        })
    return findings


def detect_underage_dob(df: pd.DataFrame, data_contract: dict, min_age: int = 18) -> list[dict]:
    dob_col = data_contract.get("dob_column")
    if not dob_col or dob_col not in df.columns:
        return []
    today = pd.Timestamp(datetime.now().date())
    ages = df[dob_col].apply(_parse_dob)
    valid = ages.dropna()
    age_years = valid.apply(lambda d: (today - d).days / 365.25)
    idx = age_years.index[age_years < min_age]
    if len(idx) == 0:
        return []
    return [{
        "check_name": "underage_dob",
        "dimension": "Validity",
        "columns": [dob_col],
        "row_refs": [{"row": int(i) + 2, "column": dob_col} for i in idx],
        "examples": [str(valid.loc[i].date()) for i in list(idx)[:5]],
        "count": int(len(idx)),
        "evidence": {"note": f"computed age from '{dob_col}' is under {min_age}"},
    }]


def detect_bad_status_transitions(df: pd.DataFrame, data_contract: dict,
                                   valid_transitions: dict = _DEFAULT_VALID_TRANSITIONS) -> list[dict]:
    status_col = data_contract.get("status_column")
    id_col = data_contract.get("id_column")
    date_col = data_contract.get("date_column")
    if not status_col or status_col not in df.columns or not id_col or id_col not in df.columns:
        return []
    work = df[[id_col, status_col] + ([date_col] if date_col and date_col in df.columns else [])].copy()
    work["__row__"] = work.index
    if date_col and date_col in work.columns:
        work["__d__"] = pd.to_datetime(work[date_col], errors="coerce", dayfirst=True)
        work = work.sort_values([id_col, "__d__"])
    bad_rows = []
    for _, group in work.groupby(id_col):
        statuses = group[status_col].astype(str).str.strip().str.lower().tolist()
        rows = group["__row__"].tolist()
        for i in range(1, len(statuses)):
            prev, curr = statuses[i - 1], statuses[i]
            allowed = valid_transitions.get(prev)
            if allowed is not None and curr not in allowed:
                bad_rows.append(int(rows[i]))
    if not bad_rows:
        return []
    return [{
        "check_name": "bad_status_transition",
        "dimension": "Consistency",
        "columns": [status_col],
        "row_refs": [{"row": r + 2, "column": status_col} for r in bad_rows],
        "examples": [],
        "count": len(bad_rows),
        "evidence": {"note": f"status change not in the allowed transition table, ordered by '{id_col}'"
                              + (f" and '{date_col}'" if date_col else "")},
    }]


def detect_orange_formatting(path: Path, sheet: str, columns: list[str]) -> list[dict]:
    """Ported from the old data_cleaning.py: flags cells/rows highlighted orange,
    a convention used in the source workbooks to mark manually-flagged rows."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return []
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {name: i for i, name in enumerate(header) if name in columns}
    if not col_idx:
        return []
    orange_row_refs = []
    for row in ws.iter_rows(min_row=2):
        row_has_orange = False
        for col_name, idx in col_idx.items():
            if idx < len(row) and _is_orange(_cell_fill_rgb(row[idx])):
                row_has_orange = True
                break
        if row_has_orange:
            orange_row_refs.append({"row": row[0].row, "column": None})
    if not orange_row_refs:
        return []
    return [{
        "check_name": "orange_formatting",
        "dimension": "Data Lineage",
        "columns": list(col_idx.keys()),
        "row_refs": orange_row_refs,
        "examples": [],
        "count": len(orange_row_refs),
        "evidence": {"note": "rows with orange cell highlighting in the source workbook (a manual flag convention)"},
    }]


GENERAL_CHECKS = [
    detect_missing_values,
    detect_duplicate_rows,
    detect_duplicate_ids,
    detect_type_inconsistency,
    detect_whitespace_case_variants,
    detect_near_spelling_variants,
    detect_outliers,
    detect_negative_zero_amounts,
    detect_invalid_future_dates,
]

CONTRACT_CHECKS = [
    detect_date_pair_order,
    detect_missing_mandatory,
    detect_underage_dob,
    detect_bad_status_transitions,
]


def run_all_detections(df: pd.DataFrame, columns: list[str], data_contract: dict,
                        workbook_path: Optional[Path] = None, sheet: Optional[str] = None,
                        file_type: str = "xlsx") -> list[dict]:
    findings: list[dict] = []
    for check in GENERAL_CHECKS:
        findings.extend(check(df, columns))
    for check in CONTRACT_CHECKS:
        findings.extend(check(df, data_contract))
    if file_type == "xlsx" and workbook_path is not None and sheet is not None:
        findings.extend(detect_orange_formatting(workbook_path, sheet, columns))
    return findings


# ─────────────────────────────────────────────
# Historical search (exact + near match via difflib, no embeddings)
# ─────────────────────────────────────────────

def build_historical_index_rows(df: pd.DataFrame, sheet: str) -> list[tuple[str, str, str, str, int]]:
    """Returns (sheet, column, value, normalized_value, count) rows ready for
    cleaning_store.add_historical_values, for every object/text column."""
    rows = []
    for col in df.columns:
        if not _is_text_column(df[col]):
            continue
        vc = df[col].dropna().astype(str).value_counts()
        for value, count in vc.items():
            rows.append((sheet, str(col), value, _normalize(value), int(count)))
    return rows


def near_match(value: str, candidates: list[dict], cutoff: float = 0.75, limit: int = 5) -> list[dict]:
    norm = _normalize(value)
    by_norm = {c["normalized_value"]: c for c in candidates}
    matches = difflib.get_close_matches(norm, list(by_norm.keys()), n=limit, cutoff=cutoff)
    return [by_norm[m] for m in matches]
